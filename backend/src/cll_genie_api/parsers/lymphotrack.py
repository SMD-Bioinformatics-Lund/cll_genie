from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REQUIRED_COLUMNS = {
    "Rank",
    "Sequence",
    "Merge count",
    "% total reads",
    "In-frame (Y/N)",
    "No Stop codon (Y/N)",
}

NUMERIC_COLUMNS = {
    "Rank": int,
    "Length": int,
    "Merge count": int,
    "% total reads": float,
    "Cumulative %": float,
    "Mutation rate to partial V-gene (%)": float,
    "V-coverage": float,
}


class LymphotrackParseError(ValueError):
    pass


def _number(value: Any, converter):
    if value is None or value == "":
        return None
    normalized = str(value).strip().replace(",", ".")
    return converter(float(normalized)) if converter is int else converter(normalized)


def parse_workbook(
    path: Path,
    *,
    sheet_name: str = "Merged Read Summary",
    header_row: int = 4,
    minimum_reads_percent: float = 0,
    in_frame: str = "B",
    no_stop_codon: str = "B",
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise LymphotrackParseError("Only .xlsx and .xlsm workbooks are supported")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise LymphotrackParseError("The LymphoTrack workbook could not be opened") from exc
    if sheet_name not in workbook.sheetnames:
        raise LymphotrackParseError(f"Worksheet '{sheet_name}' does not exist")
    sheet = workbook[sheet_name]

    excel_header_row = header_row + 1
    metadata: dict[str, Any] = {}
    for row in sheet.iter_rows(min_row=1, max_row=max(1, excel_header_row - 1), values_only=True):
        if len(row) >= 2 and row[0] not in (None, ""):
            key = str(row[0]).strip().replace(".", "_")
            metadata[key] = row[1]

    raw_headers = next(
        sheet.iter_rows(min_row=excel_header_row, max_row=excel_header_row, values_only=True)
    )
    headers = [str(value).strip() if value is not None else "" for value in raw_headers]
    missing = sorted(REQUIRED_COLUMNS.difference(headers))
    if missing:
        raise LymphotrackParseError(f"Missing required columns: {', '.join(missing)}")

    sequences: list[dict[str, Any]] = []
    for row_number, raw_row in enumerate(
        sheet.iter_rows(min_row=excel_header_row + 1, values_only=True),
        start=excel_header_row + 1,
    ):
        values = dict(zip(headers, raw_row, strict=False))
        if not values.get("Sequence"):
            continue
        try:
            for column, converter in NUMERIC_COLUMNS.items():
                if column in values:
                    values[column] = _number(values[column], converter)
        except (TypeError, ValueError) as exc:
            raise LymphotrackParseError(f"Invalid numeric value on row {row_number}") from exc

        reads_percent = float(values["% total reads"] or 0)
        if reads_percent < minimum_reads_percent:
            continue
        if in_frame != "B" and values.get("In-frame (Y/N)") != in_frame:
            continue
        if no_stop_codon != "B" and values.get("No Stop codon (Y/N)") != no_stop_codon:
            continue

        rank = int(values["Rank"])
        sequences.append(
            {
                "sequence_id": f"Seq{rank}",
                "rank": rank,
                "sequence": str(values["Sequence"]).strip(),
                "merge_count": int(values["Merge count"] or 0),
                "total_reads_percent": round(reads_percent, 4),
                "in_frame": values.get("In-frame (Y/N)") == "Y",
                "no_stop_codon": values.get("No Stop codon (Y/N)") == "Y",
                "source_row": row_number,
                "length": values.get("Length"),
                "v_gene": values.get("V-gene"),
            }
        )
    workbook.close()
    return sequences, metadata


def parse_qc(path: Path) -> dict[str, int | float]:
    values: dict[str, str] = {}
    try:
        with path.open(encoding="utf-8") as handle:
            for line in handle:
                key, value = line.rstrip("\n").split("\t", 1)
                values[key.strip()] = value.strip().replace(",", ".")
        return {
            "total_bases": int(values["totalCount"]),
            "q30_bases": int(values["countQ30"]),
            "q30_per": round(float(values["indexQ30"]), 2),
        }
    except (OSError, KeyError, ValueError) as exc:
        raise LymphotrackParseError("The LymphoTrack QC file is invalid") from exc
